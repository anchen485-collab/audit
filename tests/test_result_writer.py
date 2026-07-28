from openpyxl import load_workbook

from app.core.models import AuditResult
from app.excel.result_writer import write_audit_result_excel


def test_result_writer_repairs_mojibake_external_evidence_text(tmp_path):
    output_path = tmp_path / "audit_result.xlsx"
    result = AuditResult(
        row_number=2,
        employee_name="张冰冰",
        original_category="种植业",
        original_subcategory="种子种苗",
        original_stage="销售",
        original_company="安徽金泰种业有限公司",
        cleaned_company="安徽金泰种业有限公司",
        status="正确",
        confidence=90,
        business_scope=(
            "瀹夊窘閲戞嘲绉嶄笟鏈夐檺鍏徃 瀹夊窘閲戞嘲绉嶄笟鏈夐檺鍏徃 "
            "Anhui Jintai Seed Industry Co. , Ltd. 0551-63663889 "
            "瀹夊窘閲戞嘲绉嶄笟鏈夐檺鍏徃鎴愮珛浜�1996骞达紝"
            "涓撴敞浜庤荆妞掔瀛愮殑鐮斿彂涓庨攢鍞�"
        ),
        reason="大模型语义判断：官网证据显示主营辣椒种子研发与销售。",
    )

    write_audit_result_excel([result], output_path)

    wb = load_workbook(output_path, data_only=True)
    evidence_text = wb["审计明细"]["N2"].value

    assert "安徽金泰种业有限公司" in evidence_text
    assert "瀹夊窘" not in evidence_text


def test_result_writer_keeps_english_evidence_with_latin_accents(tmp_path):
    output_path = tmp_path / "audit_result.xlsx"
    result = AuditResult(
        row_number=2,
        employee_name="Alex",
        original_category="信息科技",
        original_subcategory="软件服务",
        original_stage="",
        original_company="Example Ltd.",
        cleaned_company="Example Ltd.",
        status="正确",
        confidence=85,
        business_scope="Café technology provides façade design services and résumé data management solutions.",
        reason="English website evidence supports the classification.",
    )

    write_audit_result_excel([result], output_path)

    wb = load_workbook(output_path, data_only=True)
    evidence_text = wb["审计明细"]["N2"].value

    assert "Café technology" in evidence_text
    assert "façade design" in evidence_text
    assert "编码异常" not in evidence_text


def test_result_writer_repairs_spanish_text_decoded_as_gb18030_mojibake(tmp_path):
    output_path = tmp_path / "audit_result.xlsx"
    result = AuditResult(
        row_number=2,
        employee_name="Alex",
        original_category="渔业",
        original_subcategory="捕捞",
        original_stage="",
        original_company="Inpesca Group",
        cleaned_company="Inpesca Group",
        status="正确",
        confidence=85,
        business_scope=(
            "Inpesca Group Direcci贸n Bo Mendekano Auzoa Grupo Inpesca Tradici贸n familiar "
            "Preparados para llegar m谩s lejos con diferentes caracter铆sticas. "
            "Europa Asia 脕frica Am茅rica Toneladas de at煤n Pa铆ses Tecnolog铆a "
            "Innovaci贸n constante a -18掳C. S茅 parte de la Compa帽铆a Internacional de Pesca."
        ),
        reason="Spanish website evidence supports the classification.",
    )

    write_audit_result_excel([result], output_path)

    wb = load_workbook(output_path, data_only=True)
    evidence_text = wb["审计明细"]["N2"].value

    assert "Dirección" in evidence_text
    assert "Tradición" in evidence_text
    assert "más lejos" in evidence_text
    assert "características" in evidence_text
    assert "Tecnología" in evidence_text
    assert "Compañía Internacional de Pesca" in evidence_text
    assert "Direcci贸n" not in evidence_text
    assert "编码异常" not in evidence_text
