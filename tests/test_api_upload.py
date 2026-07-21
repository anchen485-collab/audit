from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app


def build_employee_file(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["日期", "姓名", "一级分类", "细分", "企业名称&官网", "环节"])
    ws.append(["7月20日", "张硕", "种植业", "水果作物", "金川县雪梨果业开发有限责任公司", "梨（销售）"])
    wb.save(path)


def build_category_file(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["一级品类", "二级品类", "三级品类", "模块名称", "子模块列表"])
    ws.append(["农业", "种植业", "水果作物", "类型", "梨、苹果、水果"])
    ws.append([None, None, None, "环节", "繁育、加工、销售"])
    wb.save(path)


def test_upload_two_excels_returns_download_link(tmp_path, monkeypatch):
    monkeypatch.setenv("AUDIT_STORAGE_DIR", str(tmp_path / "storage"))
    employee_file = tmp_path / "employee.xlsx"
    category_file = tmp_path / "category.xlsx"
    build_employee_file(employee_file)
    build_category_file(category_file)

    client = TestClient(app)
    with employee_file.open("rb") as employee_fp, category_file.open("rb") as category_fp:
        response = client.post(
            "/audit/upload",
            files={
                "employee_file": ("employee.xlsx", employee_fp, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                "category_file": ("category.xlsx", category_fp, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            },
        )

    assert response.status_code == 200
    assert "/audit/download/" in response.text


def test_upload_with_invalid_employee_headers_returns_error_page(tmp_path, monkeypatch):
    monkeypatch.setenv("AUDIT_STORAGE_DIR", str(tmp_path / "storage"))
    employee_file = tmp_path / "employee.xlsx"
    category_file = tmp_path / "category.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.append(["错误表头"])
    ws.append(["金川县雪梨果业开发有限责任公司"])
    wb.save(employee_file)
    build_category_file(category_file)

    client = TestClient(app)
    with employee_file.open("rb") as employee_fp, category_file.open("rb") as category_fp:
        response = client.post(
            "/audit/upload",
            files={
                "employee_file": ("employee.xlsx", employee_fp, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                "category_file": ("category.xlsx", category_fp, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            },
        )

    assert response.status_code == 400
    assert "文件格式不符合模板" in response.text
    assert "员工录入表缺少表头" in response.text
