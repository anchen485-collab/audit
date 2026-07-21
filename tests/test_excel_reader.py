from pathlib import Path

from openpyxl import Workbook

from app.excel.input_reader import read_employee_excel


def build_employee_file(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["日期", "姓名", "一级分类", "细分", "企业名称&官网", "环节"])
    ws.append(["7月20日", "张硕", "种植业", "水果作物", "金川县雪梨果业开发有限责任公司 https://example.com", "梨（繁育、加工、销售）"])
    ws.append([None, None, None, None, None, None])
    ws.append(["7月20日", "相阳", "种植业", "", "黑龙江省佳莲种业有限公司", "稻谷（繁育）"])
    wb.save(path)


def test_read_employee_excel_keeps_rows_and_skips_empty_rows(tmp_path):
    file_path = tmp_path / "employee.xlsx"
    build_employee_file(file_path)

    records = read_employee_excel(file_path)

    assert len(records) == 2
    assert records[0].row_number == 2
    assert records[0].name == "张硕"
    assert records[0].company_raw.startswith("金川县雪梨果业")
    assert records[1].row_number == 4
    assert records[1].subcategory == ""


def test_read_employee_excel_reads_company_website_hyperlink(tmp_path):
    file_path = tmp_path / "employee_with_link.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["日期", "姓名", "一级分类", "细分", "企业名称&官网", "环节"])
    ws.append(["7月21日", "张冰冰", "种植业", "蔬菜作物", "陕西汇生源生态农业有限公司", "加工、销售"])
    ws["E2"].hyperlink = "https://www.example-agri.com"
    wb.save(file_path)

    records = read_employee_excel(file_path)

    assert records[0].website_url == "https://www.example-agri.com"


def test_read_employee_excel_allows_missing_stage_header(tmp_path):
    file_path = tmp_path / "employee_without_stage.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["日期", "姓名", "一级分类", "细分", "企业名称&官网"])
    ws.append(["7月21日", "张冰冰", "种植业", "水果作物", "金川县雪梨果业开发有限责任公司"])
    wb.save(file_path)

    records = read_employee_excel(file_path)

    assert len(records) == 1
    assert records[0].stage == ""
