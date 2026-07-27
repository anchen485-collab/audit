from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.audit.graph import run_audit_workflow
from app.company.mock_provider import MockCompanyInfoProvider


def build_employee_file(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["日期", "姓名", "一级分类", "细分", "企业名称&官网", "环节"])
    ws.append(["7月20日", "张硕", "种植业", "水果作物", "金川县雪梨果业开发有限责任公司", "梨（销售）"])
    wb.save(path)


def build_category_file(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["一级品类", "二级品类", "三级品类", "模块名称", "子模块列表"])
    ws.append(["农业", "种植业", "水果作物", "类型", "梨、苹果、水果"])
    ws.append([None, None, None, "环节", "繁育、加工、销售"])
    wb.save(path)


def test_run_audit_workflow_exports_result_excel(tmp_path, monkeypatch):
    monkeypatch.setenv("AUDIT_LLM_AGENT_ENABLED", "false")
    employee_file = tmp_path / "employee.xlsx"
    category_file = tmp_path / "category.xlsx"
    output_dir = tmp_path / "outputs"
    build_employee_file(employee_file)
    build_category_file(category_file)

    provider = MockCompanyInfoProvider(
        {
            "金川县雪梨果业开发有限责任公司": {
                "business_scope": "梨、水果种植、加工、销售。",
                "status": "存续",
            }
        }
    )

    state = run_audit_workflow(employee_file, category_file, output_dir, provider)

    assert state["output_path"]
    assert state["trace_id"].startswith("audit_")
    assert [item["stage"] for item in state["trace"]] == [
        "read_employee",
        "read_category",
        "query_company_info",
        "audit_records",
        "build_summary",
        "export_result",
        "audit_workflow",
    ]
    assert all(item["duration_ms"] >= 0 for item in state["trace"])
    assert Path(state["output_path"]).exists()
    wb = load_workbook(state["output_path"], data_only=True)
    assert "审计明细" in wb.sheetnames
    assert "人员汇总" in wb.sheetnames
    headers = [cell.value for cell in wb["审计明细"][1]]
    assert "数据来源" in headers
    assert "官网链接" in headers
    assert "外部证据文本" in headers


def test_run_audit_workflow_loads_category_rules_from_config(tmp_path, monkeypatch):
    monkeypatch.setenv("AUDIT_LLM_AGENT_ENABLED", "false")
    employee_file = tmp_path / "employee.xlsx"
    category_file = tmp_path / "category.xlsx"
    output_dir = tmp_path / "outputs"
    json_path = tmp_path / "category_rules.json"
    build_employee_file(employee_file)
    build_category_file(category_file)
    monkeypatch.setenv("CATEGORY_RULES_EXCEL_PATH", str(category_file))
    monkeypatch.setenv("CATEGORY_RULES_JSON_PATH", str(json_path))

    provider = MockCompanyInfoProvider(
        {
            "金川县雪梨果业开发有限责任公司": {
                "business_scope": "梨、水果种植、加工、销售。",
                "status": "存续",
            }
        }
    )

    state = run_audit_workflow(employee_file=employee_file, output_dir=output_dir, provider=provider)

    assert json_path.exists()
    assert state["output_path"]
    assert Path(state["output_path"]).exists()
