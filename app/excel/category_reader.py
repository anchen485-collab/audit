from pathlib import Path
import re

from openpyxl import load_workbook

from app.core.models import CategoryRule
from app.excel.input_reader import cell_text


REQUIRED_HEADERS = ["一级品类", "二级品类", "三级品类", "模块名称", "子模块列表"]


def split_keywords(text: str) -> list[str]:
    """按常见中英文分隔符拆分子模块关键词。"""
    parts = re.split(r"[、，,;；\n\r]+", text)
    return [part.strip() for part in parts if part and part.strip()]


def read_category_rules(path: str | Path) -> list[CategoryRule]:
    """读取内部分类表，并处理合并单元格展开后的空值继承。"""
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    header_row = [cell_text(cell.value) for cell in ws[1]]
    header_index = {name: header_row.index(name) for name in REQUIRED_HEADERS if name in header_row}
    missing_headers = [name for name in REQUIRED_HEADERS if name not in header_index]
    if missing_headers:
        raise ValueError(f"内部分类表缺少表头：{', '.join(missing_headers)}")

    current_level1 = ""
    current_level2 = ""
    current_level3 = ""
    rules: list[CategoryRule] = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        values = [cell_text(value) for value in row]
        level1 = values[header_index["一级品类"]]
        level2 = values[header_index["二级品类"]]
        level3 = values[header_index["三级品类"]]
        module_name = values[header_index["模块名称"]]
        keyword_text = values[header_index["子模块列表"]]

        # 分类表常见合并单元格，openpyxl 读取时下方行会变空，这里继承上方值。
        if level1:
            current_level1 = level1
        if level2:
            current_level2 = level2
        if level3:
            current_level3 = level3

        if not module_name and not keyword_text:
            continue

        rules.append(
            CategoryRule(
                level1=current_level1,
                level2=current_level2,
                level3=current_level3,
                module_name=module_name,
                keywords=split_keywords(keyword_text),
            )
        )
    return rules
