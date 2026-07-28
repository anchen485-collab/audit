from pathlib import Path

from openpyxl import load_workbook

from app.core.models import EmployeeRecord


REQUIRED_HEADER_ALIASES = {
    "日期": ["日期"],
    "姓名": ["姓名"],
    "一级分类": ["一级分类"],
    "细分": ["细分", "细项"],
    "企业名称&官网": ["企业名称&官网"],
}


def cell_text(value) -> str:
    """把 Excel 单元格值统一转成去空格字符串。"""
    if value is None:
        return ""
    return str(value).strip()


def read_employee_excel(path: str | Path) -> list[EmployeeRecord]:
    """读取员工录入表，保留原始行号并跳过完全空行。"""
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    header_row = [cell_text(cell.value) for cell in ws[1]]
    header_index = _resolve_header_index(header_row)
    missing_headers = ["/".join(aliases) for name, aliases in REQUIRED_HEADER_ALIASES.items() if name not in header_index]
    if missing_headers:
        raise ValueError(f"员工录入表缺少表头：{', '.join(missing_headers)}")

    records: list[EmployeeRecord] = []
    company_col = header_index["企业名称&官网"]
    for row_number, row in enumerate(ws.iter_rows(min_row=2), start=2):
        values = [cell_text(cell.value) for cell in row]
        if not any(values):
            continue
        company_cell = row[company_col]
        website_url = ""
        if company_cell.hyperlink and company_cell.hyperlink.target:
            website_url = str(company_cell.hyperlink.target).strip()
        records.append(
            EmployeeRecord(
                row_number=row_number,
                date=values[header_index["日期"]],
                name=values[header_index["姓名"]],
                category=values[header_index["一级分类"]],
                subcategory=values[header_index["细分"]],
                company_raw=values[header_index["企业名称&官网"]],
                website_url=website_url,
            )
        )
    return records


def _resolve_header_index(header_row: list[str]) -> dict[str, int]:
    header_index: dict[str, int] = {}
    for standard_name, aliases in REQUIRED_HEADER_ALIASES.items():
        for alias in aliases:
            if alias in header_row:
                header_index[standard_name] = header_row.index(alias)
                break
    return header_index
