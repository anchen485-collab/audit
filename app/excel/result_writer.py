from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.audit.summary import build_person_summary
from app.company.text_repair import looks_mojibake, repair_compacted_text
from app.core.models import AuditResult


DETAIL_HEADERS = [
    "原始行号",
    "姓名",
    "一级分类",
    "细分",
    "环节",
    "企业名称原文",
    "清洗后企业名称",
    "审计结果",
    "置信度",
    "数据来源",
    "官网链接",
    "企查查企业名称",
    "经营状态",
    "外部证据文本",
    "错误类型",
    "错误原因",
    "建议修正",
    "是否需要人工复核",
]


STATUS_FILL = {
    "正确": "D9EAD3",
    "错误": "F4CCCC",
    "疑似错误": "FCE5CD",
    "无法判断": "D9EAF7",
    "信息缺失": "FFF2CC",
}


def write_audit_result_excel(results: list[AuditResult], output_path: str | Path) -> Path:
    """导出审计明细和人员汇总两个 Sheet。"""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    detail_ws = wb.active
    detail_ws.title = "审计明细"
    detail_ws.append(DETAIL_HEADERS)

    for result in results:
        detail_ws.append(
            [
                result.row_number,
                result.employee_name,
                result.original_category,
                result.original_subcategory,
                result.original_stage,
                result.original_company,
                result.cleaned_company,
                result.status,
                result.confidence,
                result.data_source,
                result.website_url,
                result.company_name,
                result.company_status,
                _external_evidence_for_display(result.business_scope),
                result.error_type,
                result.reason,
                result.suggestion,
                "是" if result.needs_review else "否",
            ]
        )

    summary_ws = wb.create_sheet("人员汇总")
    summary_rows = build_person_summary(results)
    summary_headers = ["姓名", "录入数量", "完整数量", "正确数量", "错误数量", "疑似错误数量", "无法判断数量", "正确率"]
    summary_ws.append(summary_headers)
    for item in summary_rows:
        summary_ws.append([item[header] for header in summary_headers])

    _style_sheet(detail_ws)
    _style_sheet(summary_ws)
    wb.save(output_path)
    return output_path


def _style_sheet(ws):
    """添加基础样式，方便业务人员打开 Excel 后直接查看。"""
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    ws.freeze_panes = "A2"
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 10), 42)
    if ws.title == "审计明细":
        for row in ws.iter_rows(min_row=2):
            status = row[7].value
            fill_color = STATUS_FILL.get(status)
            if fill_color:
                row[7].fill = PatternFill("solid", fgColor=fill_color)


def _external_evidence_for_display(text: str) -> str:
    """导出前再做一次展示文本修复，避免 Excel 明细里残留爬虫乱码。"""
    repaired = repair_compacted_text(text)
    if looks_mojibake(repaired):
        return "外部证据文本编码异常，已隐藏乱码；请清理官网缓存后重新审计"
    return repaired
